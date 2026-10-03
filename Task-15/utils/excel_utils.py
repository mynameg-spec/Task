"""
excel_utils.py
Reusable helper class to read test data from and write results to the Excel file.
This is the "Data Driven" part of the framework.
"""
import os
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.exceptions import InvalidFileException


class ExcelUtils:
    """Read and write login test data in an Excel (.xlsx) file."""

    # Column positions in the Excel sheet (1-based)
    HEADERS = ["Test ID", "Username", "Password", "Date", "Time of Test", "Name of Tester", "Test Result"]
    COL_TEST_ID = 1
    COL_USERNAME = 2
    COL_PASSWORD = 3
    COL_DATE = 4
    COL_TIME = 5
    COL_TESTER = 6
    COL_RESULT = 7

    def __init__(self, file_path, sheet_name):
        self.file_path = file_path
        self.sheet_name = sheet_name

    def _load_workbook(self):
        """Open the workbook, raising a clear message if it is missing or invalid."""
        try:
            return load_workbook(self.file_path)
        except FileNotFoundError:
            raise FileNotFoundError(f"Excel file not found: {self.file_path}")
        except InvalidFileException:
            raise InvalidFileException(f"Not a valid .xlsx file: {self.file_path}")

    def _get_sheet(self, workbook):
        """Return the required sheet from the workbook."""
        if self.sheet_name not in workbook.sheetnames:
            raise KeyError(f"Sheet '{self.sheet_name}' not found in {self.file_path}")
        return workbook[self.sheet_name]

    def get_row_count(self):
        """Return the total number of rows (including the header row)."""
        workbook = self._load_workbook()
        row_count = self._get_sheet(workbook).max_row
        workbook.close()
        return row_count

    def read_cell(self, row, column):
        """Return the value of a single cell."""
        workbook = self._load_workbook()
        value = self._get_sheet(workbook).cell(row=row, column=column).value
        workbook.close()
        return value

    def write_cell(self, row, column, value):
        """Write a value into a single cell and save the file."""
        workbook = self._load_workbook()
        self._get_sheet(workbook).cell(row=row, column=column, value=value)
        self._save(workbook)

    def get_login_test_data(self):
        """
        Read all data rows and return them as a list of dictionaries.
        Each dictionary also stores its Excel row number so results can be
        written back to the same row.
        """
        workbook = self._load_workbook()
        sheet = self._get_sheet(workbook)
        test_data = []

        for row in range(2, sheet.max_row + 1):  # Row 1 is the header
            test_id = sheet.cell(row=row, column=self.COL_TEST_ID).value
            if test_id is None:
                continue  # Skip empty rows
            test_data.append({
                "row": row,
                "test_id": str(test_id),
                "username": str(sheet.cell(row=row, column=self.COL_USERNAME).value or ""),
                "password": str(sheet.cell(row=row, column=self.COL_PASSWORD).value or ""),
            })

        workbook.close()
        return test_data

    def write_test_result(self, row, tester_name, result):
        """Write Date, Time of Test, Name of Tester and Test Result for one row."""
        now = datetime.now()
        workbook = self._load_workbook()
        sheet = self._get_sheet(workbook)

        sheet.cell(row=row, column=self.COL_DATE, value=now.strftime("%d-%m-%Y"))
        sheet.cell(row=row, column=self.COL_TIME, value=now.strftime("%H:%M:%S"))
        sheet.cell(row=row, column=self.COL_TESTER, value=tester_name)

        result_cell = sheet.cell(row=row, column=self.COL_RESULT, value=result)
        # Green for Passed, red for Failed - easy to read at a glance
        colour = "C6EFCE" if result == "Test Passed" else "FFC7CE"
        result_cell.fill = PatternFill(start_color=colour, end_color=colour, fill_type="solid")

        self._save(workbook)

    def _save(self, workbook):
        """Save the workbook. Fails clearly if the file is open in Excel."""
        try:
            workbook.save(self.file_path)
        except PermissionError:
            raise PermissionError(f"Close the Excel file before running tests: {self.file_path}")
        finally:
            workbook.close()

    def create_test_data_file(self, login_rows):
        """Create the Excel file with headers and the given (test_id, username, password) rows."""
        os.makedirs(os.path.dirname(self.file_path), exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self.sheet_name

        sheet.append(self.HEADERS)
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        for test_id, username, password in login_rows:
            sheet.append([test_id, username, password, None, None, None, None])

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial")

        column_widths = [10, 18, 18, 14, 14, 18, 16]
        for index, width in enumerate(column_widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width

        workbook.save(self.file_path)
        workbook.close()
